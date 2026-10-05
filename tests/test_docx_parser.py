# -*- coding: utf-8 -*-
"""End-to-end test of the DOCX parser using a document built in-memory."""

import tempfile
import unittest
from pathlib import Path

from docx import Document  # type: ignore
from docx.shared import Inches  # type: ignore

from tomd.parsers.docx_parser import DocxParser


class DocxParserTest(unittest.TestCase):
    def _build_docx(self, path: Path) -> None:
        doc = Document()
        doc.add_paragraph("A study of procurement", style="Title")
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("We examine contracting outcomes.")
        doc.add_heading("Methods", level=2)
        doc.add_paragraph("Our sample is large. We estimate a regression.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Year"
        table.cell(0, 1).text = "Effect"
        table.cell(1, 0).text = "2020"
        table.cell(1, 1).text = "0.31"
        doc.add_paragraph("Figure caption text goes here")
        doc.save(str(path))

    def test_heading_levels_and_table(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "paper.docx"
            self._build_docx(path)
            result = DocxParser().parse(str(path), "paper.docx")

            body = result.body
            self.assertEqual(result.source_kind.value, "docx")
            # Title -> one #, Heading1 -> ##, Heading2 -> ###
            self.assertIn("# A study of procurement", body)
            self.assertIn("## Introduction", body)
            self.assertIn("### Methods", body)
            # markdown table rendered with pipe separators and a header row
            self.assertIn("| Year | Effect |", body)
            self.assertIn("|---|", body)  # header delimiter row
            self.assertIn("| 2020 | 0.31 |", body)

    def test_corrupt_file_raises(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            bad = Path(td) / "bad.docx"
            bad.write_bytes(b"this is not a real docx zip")
            with self.assertRaises(Exception):
                DocxParser().parse(str(bad), "bad.docx")


if __name__ == "__main__":
    unittest.main()